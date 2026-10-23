"""Administrator export: four formats, filters like the list, no drafts, BOM and quoting, audit (QE-173 to QE-184)."""

from __future__ import annotations

import csv
import io
import json
from datetime import date

import pytest
from _feedback_support import make_env

COLUMNS = [
    "id",
    "title",
    "subject",
    "project",
    "target_kind",
    "target_id",
    "priority",
    "complexity",
    "status",
    "submitter",
    "created_at",
    "confirmed_at",
    "duplicate_of",
    "conflicts",
    "summary",
    "description",
    "business_need",
    "justification",
    "agent_priority",
    "agent_complexity",
]


@pytest.fixture(autouse=True, scope="module")
def real_openpyxl():
    """``tests/conftest.py`` replaces ``openpyxl`` with a mock; the Excel export needs the real package."""
    import importlib
    import sys
    from unittest import mock

    saved = {
        k: v
        for k, v in sys.modules.items()
        if k == "openpyxl" or k.startswith("openpyxl.")
    }
    if any(isinstance(v, mock.Mock) for v in saved.values()):
        for k in saved:
            del sys.modules[k]
        importlib.import_module("openpyxl")
    try:
        yield
    finally:
        for k in [
            k for k in sys.modules if k == "openpyxl" or k.startswith("openpyxl.")
        ]:
            del sys.modules[k]
        sys.modules.update(saved)


@pytest.fixture(params=["sqlite", "postgres"])
def env(request, monkeypatch):
    with make_env(monkeypatch, request.param) as e:
        e.confirmed(
            "omar",
            'Export, "quoted" title',
            number=1,
            priority="high",
            complexity="L",
            description="Line one\nLine two, with comma",
            conflicts=[
                {"request_id": "FB-0002", "kind": "conflict", "note": "n"},
                {"request_id": "FB-0003", "kind": "related", "note": "n"},
            ],
            agent_priority="medium",
            agent_complexity="M",
        )
        e.confirmed(
            "sara",
            "Dark mode",
            number=2,
            status="accepted",
            priority="low",
            complexity="S",
        )
        e.confirmed("sara", "Arabic: تصدير", number=3, priority="low", complexity="S")
        from qym_platform.db.feedback_models import FeedbackRequest

        with e.db() as db:
            db.add(
                FeedbackRequest(
                    subject="platform",
                    status="draft",
                    title="Secret draft",
                    submitter_user_id=e.users["admin"],
                    checklist={},
                    conflicts=[],
                    agent_reasoning={},
                )
            )
            db.commit()
        e.act_as("admin")
        yield e


def test_csv_has_bom_header_order_quoting_and_download_name(env):
    r = env.client.get("/v1/feedback/export?format=csv")
    assert r.status_code == 200
    assert r.headers["content-type"] == "text/csv; charset=utf-8"
    assert (
        r.headers["content-disposition"]
        == f'attachment; filename="feedback-{date.today().isoformat()}.csv"'
    )
    assert r.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0] == COLUMNS
    assert len(rows) == 4  # header + three confirmed, never the draft
    one = next(row for row in rows[1:] if row[0] == "FB-0001")
    rec = dict(zip(COLUMNS, one))
    assert (
        rec["title"] == 'Export, "quoted" title'
        and rec["description"] == "Line one\nLine two, with comma"
    )
    assert rec["conflicts"] == "FB-0002;FB-0003" and rec["agent_priority"] == "medium"
    assert "Secret draft" not in r.text
    assert env.client.get("/v1/feedback/export.csv").content == r.content


def test_json_jsonl_and_xlsx_carry_the_same_columns(env):
    j = env.client.get("/v1/feedback/export?format=json")
    assert j.headers["content-type"].startswith("application/json")
    data = json.loads(j.content)
    assert [list(d) for d in data][0] == COLUMNS and len(data) == 3
    first = next(d for d in data if d["id"] == "FB-0001")
    assert (
        first["conflicts"] == ["FB-0002", "FB-0003"]
        and first["project"] is None
        and first["created_at"].endswith("Z")
    )
    nd = env.client.get("/v1/feedback/export?format=jsonl")
    assert nd.headers["content-type"].startswith("application/x-ndjson")
    assert nd.headers["content-disposition"].endswith('.jsonl"')
    lines = nd.content.decode().strip().split("\n")
    assert len(lines) == 3 and json.loads(lines[0]).keys() == data[0].keys()
    x = env.client.get("/v1/feedback/export?format=xlsx")
    assert (
        x.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert x.headers["content-disposition"].endswith('.xlsx"')
    from datetime import datetime

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Feedback"]
    ws = wb["Feedback"]
    assert (
        [c.value for c in ws[1]] == COLUMNS
        and ws[1][0].font.bold
        and ws.freeze_panes == "A2"
    )
    row = next(
        r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "FB-0001"
    )
    assert row[13] == "FB-0002;FB-0003" and isinstance(row[10], datetime)


def test_filters_match_the_list_and_paging_is_ignored(env):
    for query in (
        "status=accepted",
        "priority=low",
        "q=dark",
        "complexity=L",
        "about=platform&sort=priority:desc",
    ):
        listed = [
            i["ref"]
            for i in env.client.get(f"/v1/feedback?{query}&page_size=100").json()[
                "items"
            ]
        ]
        exported = [
            d["id"]
            for d in json.loads(
                env.client.get(f"/v1/feedback/export?format=json&{query}").content
            )
        ]
        assert exported == listed, query
    assert (
        len(
            json.loads(
                env.client.get(
                    "/v1/feedback/export?format=json&page=2&page_size=1&mine=true"
                ).content
            )
        )
        == 3
    )
    assert (
        json.loads(
            env.client.get("/v1/feedback/export?format=json&status=draft").content
        )
        == []
    )


def test_members_get_403_unknown_formats_422_and_each_export_is_audited(env):
    from qym_platform.db.models import AuditLog

    bad = env.client.get("/v1/feedback/export?format=pdf")
    assert bad.status_code == 422 and bad.json()["accepted"] == [
        "csv",
        "xlsx",
        "json",
        "jsonl",
    ]
    env.act_as("omar")
    for path in ("/v1/feedback/export?format=csv", "/v1/feedback/export.csv"):
        assert env.client.get(path).status_code == 403
    env.act_as("admin")
    env.client.get("/v1/feedback/export?format=json&status=new&q=export")
    with env.db() as db:
        entries = db.query(AuditLog).filter_by(action="feedback.export").all()
        assert len(entries) == 1  # the refused requests wrote nothing
        assert entries[0].after["format"] == "json" and entries[0].after["rows"] == 1
        assert entries[0].after["filter"] == {"status": "new", "q": "export"}
        assert (
            "Export" not in str(entries[0].after["filter"])
            and "title" not in entries[0].after
        )
