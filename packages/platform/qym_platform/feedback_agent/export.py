"""Administrator export of the feedback list: CSV, Excel, JSON and JSON Lines (G6, QE-173 to QE-184)."""

from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime
from typing import Any, Iterable, Optional

from qym_platform.datetime_utils import to_api_timestamp
from qym_platform.db.feedback_models import FeedbackRequest
from qym_platform.db.models import Project, User
from qym_platform.feedback_agent.search import request_ref
from qym_platform.feedback_agent.store import user_name

COLUMNS = (
    "id",
    "title",
    "subject",
    "project",
    "target_kind",
    "target_id",
    "priority",
    "complexity",
    "status",
    "submitter",
    "created_at",
    "confirmed_at",
    "duplicate_of",
    "conflicts",
    "summary",
    "description",
    "business_need",
    "justification",
    "agent_priority",
    "agent_complexity",
)
FORMATS = {
    "csv": ("text/csv; charset=utf-8", "csv"),
    "xlsx": (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    ),
    "json": ("application/json", "json"),
    "jsonl": ("application/x-ndjson", "jsonl"),
}


def _conflict_ids(row: FeedbackRequest) -> list[str]:
    return [
        str(c.get("request_id"))
        for c in (row.conflicts or [])
        if isinstance(c, dict) and c.get("request_id")
    ]


def export_records(
    rows: Iterable[FeedbackRequest],
    users: dict[str, User],
    projects: dict[str, Project],
    refs: dict[str, str],
) -> list[dict[str, Any]]:
    """Plain records with real values (datetimes stay datetimes; ``conflicts`` is a list)."""
    out = []
    for r in rows:
        project = projects.get(r.project_id) if r.project_id else None
        out.append(
            {
                "id": request_ref(r.number),
                "title": r.title,
                "subject": r.subject,
                "project": project.name if project else None,
                "target_kind": r.target_kind,
                "target_id": r.target_id,
                "priority": r.priority,
                "complexity": r.complexity,
                "status": r.status,
                "submitter": user_name(users.get(r.submitter_user_id)),
                "created_at": r.created_at,
                "confirmed_at": r.confirmed_at,
                "duplicate_of": (
                    refs.get(r.duplicate_of_id) if r.duplicate_of_id else None
                ),
                "conflicts": _conflict_ids(r),
                "summary": r.summary,
                "description": r.description,
                "business_need": r.business_need,
                "justification": r.justification,
                "agent_priority": r.agent_priority,
                "agent_complexity": r.agent_complexity,
            }
        )
    return out


def _json_ready(record: dict[str, Any]) -> dict[str, Any]:
    out = dict(record)
    for key in ("created_at", "confirmed_at"):
        out[key] = (
            to_api_timestamp(out[key]) if isinstance(out[key], datetime) else out[key]
        )
    return out


def to_csv(records: list[dict[str, Any]]) -> bytes:
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(COLUMNS)
    for rec in records:
        r = _json_ready(rec)
        writer.writerow(
            [
                ";".join(r[c]) if c == "conflicts" else ("" if r[c] is None else r[c])
                for c in COLUMNS
            ]
        )
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")


def to_json(records: list[dict[str, Any]]) -> bytes:
    return json.dumps([_json_ready(r) for r in records], ensure_ascii=False).encode(
        "utf-8"
    )


def to_jsonl(records: list[dict[str, Any]]) -> bytes:
    return "".join(
        json.dumps(_json_ready(r), ensure_ascii=False) + "\n" for r in records
    ).encode("utf-8")


def to_xlsx(records: list[dict[str, Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "Feedback"
    ws.append(list(COLUMNS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for rec in records:
        row = []
        for c in COLUMNS:
            v = rec[c]
            if c == "conflicts":
                v = ";".join(v)
            elif isinstance(v, datetime):
                v = v.replace(tzinfo=None)
            row.append(v)
        ws.append(row)
    for idx, name in enumerate(COLUMNS, start=1):
        letter = ws.cell(row=1, column=idx).column_letter
        ws.column_dimensions[letter].width = max(12, len(name) + 4)
        if name in ("created_at", "confirmed_at"):
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=idx).number_format = "yyyy-mm-dd hh:mm:ss"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render(fmt: str, records: list[dict[str, Any]]) -> bytes:
    return {"csv": to_csv, "xlsx": to_xlsx, "json": to_json, "jsonl": to_jsonl}[fmt](
        records
    )


def filename(fmt: str, today: Optional[date] = None) -> str:
    return f"feedback-{(today or date.today()).isoformat()}.{FORMATS[fmt][1]}"
